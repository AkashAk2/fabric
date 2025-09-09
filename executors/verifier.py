from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from models.contracts import QCItem, Plan, StyleHints
from typing import List, Set

def _ratio(w,h): return round(w/h, 3)

def _is_font_ok(name: str, allowed: Set[str]):
    if not name: return False
    low = name.strip().lower()
    return any(a.lower() in low for a in allowed)

def verify_pptx(path: str, plan: Plan, style: StyleHints, guidelines: dict) -> List[QCItem]:
    prs = Presentation(path); g = guidelines; out: List[QCItem] = []

    # Aspect ratio
    want_w, want_h = g["visual_rules"]["canvas_px"]; want_ratio = _ratio(want_w, want_h)
    if abs(_ratio(prs.slide_width, prs.slide_height) - want_ratio) > 0.01:
        out.append(QCItem(criterion="Aspect ratio 16:9", result="FAIL", reason="Wrong slide size", severity="BLOCKER"))

    # Slide count
    n = len(prs.slides)
    if not (g["content_rules"]["slide_count_min"] <= n <= g["content_rules"]["slide_count_max"]):
        out.append(QCItem(criterion="Slide count", result="FAIL", reason=f"{n} not within bounds", severity="BLOCKER"))

    # Titles, bullets (count/length), fonts, sizes, image presence/alt text
    allowed_families = {"Arial","Calibri"}
    min_title = g["brand"]["fonts"]["title_min_pt"]; min_body = g["brand"]["fonts"]["body_min_pt"]
    min_bul = g["content_rules"]["bullets_per_slide_min"]; max_bul = g["content_rules"]["bullets_per_slide_max"]; max_len = g["content_rules"]["bullet_char_limit"]

    for i, slide in enumerate(prs.slides):
        # title
        if not (slide.shapes.title and slide.shapes.title.text and slide.shapes.title.text.strip()):
            out.append(QCItem(criterion=f"Slide {i+1} title", result="FAIL", reason="Missing", severity="BLOCKER"))
        else:
            tf = slide.shapes.title.text_frame
            for run in tf.paragraphs[0].runs or []:
                sz = int(run.font.size.pt) if run.font.size else min_title
                if sz < min_title: out.append(QCItem(criterion=f"Slide {i+1} title size", result="FAIL", reason=f"{sz}pt < {min_title}pt", severity="MAJOR"))
                if not _is_font_ok(run.font.name or "", allowed_families):
                    out.append(QCItem(criterion=f"Slide {i+1} title font", result="FAIL", reason=f"{run.font.name} not allowed", severity="MAJOR"))

        # bullets (best-effort)
        bullets = 0
        for sh in slide.shapes:
            if not getattr(sh, "has_text_frame", False): continue
            tf = sh.text_frame
            for p in tf.paragraphs:
                txt = (p.text or "").strip()
                if txt and p.level == 0:
                    bullets += 1
                    if len(txt) > max_len:
                        out.append(QCItem(criterion=f"Slide {i+1} bullet length", result="FAIL", reason="> 120 chars", severity="MAJOR"))
                    for r in p.runs:
                        sz = int(r.font.size.pt) if r.font.size else min_body
                        if sz < min_body:
                            out.append(QCItem(criterion=f"Slide {i+1} body size", result="FAIL", reason=f"{sz}pt < {min_body}pt", severity="MAJOR"))
                        if not _is_font_ok(r.font.name or "", allowed_families):
                            out.append(QCItem(criterion=f"Slide {i+1} body font", result="FAIL", reason=f"{r.font.name} not allowed", severity="MAJOR"))

        if bullets and not (min_bul <= bullets <= max_bul):
            out.append(QCItem(criterion=f"Slide {i+1} bullets per slide", result="FAIL", reason=f"{bullets} bullets", severity="MAJOR"))

        # images + alt text if required
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if i < len(plan.slides) and any(a.startswith("image:") for a in plan.slides[i].required_assets):
            if not pictures:
                out.append(QCItem(criterion=f"Slide {i+1} image required", result="FAIL", reason="Missing image", severity="MAJOR"))
        for pic in pictures:
            alt = getattr(pic, "alternative_text", "") or ""
            if not alt.strip():
                out.append(QCItem(criterion=f"Slide {i+1} image alt text", result="FAIL", reason="Missing alt", severity="MINOR"))

    # TODO: enforce brand palette, margins >=60px, WCAG AA contrast via color sampling; icons vector check
    # TODO: links_work check if hyperlinks exist

    return out

def extract_ppt_facts(path: str) -> dict:
    prs = Presentation(path)
    facts = {"slides":[]}
    for i, s in enumerate(prs.slides):
        pics = [x for x in s.shapes if x.shape_type == MSO_SHAPE_TYPE.PICTURE]
        facts["slides"].append({"index": i, "title": (s.shapes.title.text.strip() if s.shapes.title else ""), "picture_count": len(pics)})
    return facts
